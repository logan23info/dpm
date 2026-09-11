#!/usr/bin/env python3
"""
Generate a fieldwork-ready RCM workbook from the built control library.

Usage:
    node scripts/build-library.js          # regenerate data/frameworks/*.json first
    python3 scripts/export_xlsx.py [out.xlsx]

Produces three sheets:
    Legend      how to use the workbook, which columns to fill, one example row
    RCM         the control matrix; frozen headers, autofilter, dropdowns on fieldwork columns
    Summary     live COUNTIFS rollups by framework, domain, risk, and test result

Requires openpyxl.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
FRAMEWORK_DIR = ROOT / "data" / "frameworks"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "rcm-workpaper.xlsx"

FONT = "Arial"

# --- palette -----------------------------------------------------------------
INK = "12213A"
HEADER_FILL = PatternFill("solid", fgColor=INK)
GROUP_FILLS = {
    "Identification": PatternFill("solid", fgColor="D9DDE5"),
    "Control": PatternFill("solid", fgColor="E4E9F0"),
    "Testing approach": PatternFill("solid", fgColor="E7EFE9"),
    "Fieldwork (complete during testing)": PatternFill("solid", fgColor="FFF2CC"),
}
INPUT_FILL = PatternFill("solid", fgColor="FFFDE7")  # cells the auditor fills in
RISK_FILLS = {
    "High": PatternFill("solid", fgColor="F6E3E0"),
    "Medium": PatternFill("solid", fgColor="F3E7D2"),
    "Low": PatternFill("solid", fgColor="E1EFE7"),
}
RISK_FONTS = {
    "High": Font(name=FONT, size=10, color="A23E33", bold=True),
    "Medium": Font(name=FONT, size=10, color="8A6220", bold=True),
    "Low": Font(name=FONT, size=10, color="3F7A5C", bold=True),
}
THIN = Side(style="thin", color="C7CDD8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- column layout -----------------------------------------------------------
# (field, header, width, group)
COLUMNS = [
    ("id", "Control ID", 20, "Identification"),
    ("framework", "Framework", 12, "Identification"),
    ("clause_ref", "Clause Ref", 20, "Identification"),
    ("domain", "Domain", 30, "Identification"),

    ("control_objective", "Control Objective", 50, "Control"),
    ("control_description", "Control Description", 60, "Control"),
    ("risk_addressed", "Risk Addressed", 50, "Control"),
    ("inherent_risk", "Inherent Risk", 13, "Control"),
    ("key_control", "Key Control", 11, "Control"),
    ("control_type", "Control Type", 13, "Control"),
    ("control_nature", "Control Nature", 20, "Control"),
    ("frequency", "Frequency", 14, "Control"),
    ("owner", "Control Owner", 24, "Control"),
    ("itgc_dependency", "ITGC Dependencies", 26, "Control"),
    ("cross_framework_refs", "Cross-Framework Mapping", 34, "Control"),

    ("test_of_design", "Test of Design", 55, "Testing approach"),
    ("test_of_effectiveness", "Test of Effectiveness", 55, "Testing approach"),
    ("population_basis", "Population Basis", 45, "Testing approach"),
    ("sampling_basis", "Sampling Basis", 45, "Testing approach"),
    ("evidence_required", "Evidence Required", 45, "Testing approach"),

    ("status", "Implementation Status", 20, "Fieldwork (complete during testing)"),
    ("test_result", "Test Result", 24, "Fieldwork (complete during testing)"),
    ("residual_risk", "Residual Risk", 14, "Fieldwork (complete during testing)"),
    ("exceptions_noted", "Exceptions Noted", 40, "Fieldwork (complete during testing)"),
    ("conclusion", "Conclusion", 40, "Fieldwork (complete during testing)"),
    ("gap_notes", "Gap Notes", 34, "Fieldwork (complete during testing)"),
    ("period_covered", "Period Covered", 22, "Fieldwork (complete during testing)"),
    ("prepared_by", "Prepared By", 18, "Fieldwork (complete during testing)"),
    ("reviewed_by", "Reviewed By", 18, "Fieldwork (complete during testing)"),
    ("last_reviewed", "Date Reviewed", 15, "Fieldwork (complete during testing)"),
]

FIELDWORK_GROUP = "Fieldwork (complete during testing)"

DROPDOWNS = {
    "status": ["Not Started", "In Progress", "Implemented", "Not Applicable"],
    "test_result": ["Not Tested", "Effective", "Effective with Exceptions", "Ineffective"],
    "residual_risk": ["High", "Medium", "Low"],
}

EXAMPLE_ROW = {
    "status": "Implemented",
    "test_result": "Effective with Exceptions",
    "residual_risk": "Medium",
    "exceptions_noted": "2 of 25 DSARs closed at 34 and 41 days against a 30-day SLA.",
    "conclusion": "Control operating, exceptions isolated to Q2 staffing gap; remediated.",
    "gap_notes": "Escalation alert not configured until 12 May.",
    "period_covered": "2026-01-01 to 2026-06-30",
    "prepared_by": "A. Auditor",
    "reviewed_by": "S. Manager",
    "last_reviewed": "2026-07-15",
}


def load_controls():
    controls = []
    if not FRAMEWORK_DIR.exists():
        sys.exit(
            "data/frameworks not found. Run 'node scripts/build-library.js' first."
        )
    for path in sorted(FRAMEWORK_DIR.glob("*.json")):
        controls.extend(json.loads(path.read_text()))
    # Stable, review-friendly order: framework, then domain, then ID.
    controls.sort(key=lambda c: (c["framework"], c["domain"], c["id"]))
    return controls


def cell_value(control, field):
    v = control.get(field)
    if v is None:
        return None
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, list):
        return "; ".join(v) if v else None
    return v


def build_legend(ws, control_count, generated_from):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95

    ws["A1"] = "RCM Studio — Privacy & Data Protection Control Matrix"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=INK)
    ws.merge_cells("A1:B1")

    rows = [
        ("", ""),
        ("Purpose", "Multi-framework Risk Control Matrix for privacy and data protection audits. "
                    "Covers ISO 27701, GDPR, India's DPDP Act 2023, NIST Privacy Framework and SOC 2 privacy criteria."),
        ("Controls in this workbook", control_count),
        ("Generated from", generated_from),
        ("", ""),
        ("HOW TO USE", ""),
        ("1. Scope", "On the RCM sheet, use the header filters to select your framework(s). "
                     "Filter Key Control = Yes and Inherent Risk = High for a defensible key-control population."),
        ("2. Test", "Perform the procedures in Test of Design and Test of Effectiveness. "
                    "Sampling Basis states the sample size and how it is justified by control frequency."),
        ("3. Record", "Complete the shaded Fieldwork columns. Those cells are the only ones you edit."),
        ("4. Roll up", "The Summary sheet updates automatically from what you enter — do not overwrite it."),
        ("", ""),
        ("WHICH CELLS TO EDIT", ""),
        ("Editable", "The pale yellow Fieldwork columns only: Implementation Status, Test Result, Residual Risk, "
                     "Exceptions Noted, Conclusion, Gap Notes, Period Covered, Prepared By, Reviewed By, Date Reviewed."),
        ("Do not edit", "Identification, Control and Testing Approach columns. These are generated from the "
                        "control library in the source repository. Change them at source and re-export, "
                        "so the library stays the single point of truth."),
        ("Dropdowns", "Implementation Status, Test Result and Residual Risk are restricted to valid values."),
        ("", ""),
        ("EXAMPLE ROW", "Realistic values showing the expected format for the Fieldwork columns:"),
    ]

    r = 2
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(
            name=FONT, size=10, bold=label.isupper() or label in ("Purpose", "Controls in this workbook", "Generated from")
        )
        c = ws.cell(row=r, column=2, value=value)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if label.isupper() and label:
            ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True, color=INK)
        r += 1

    # Example row, laid out as label/value pairs so it cannot be mistaken for data.
    r += 1
    for field, value in EXAMPLE_ROW.items():
        header = next(h for f, h, _, _ in COLUMNS if f == field)
        lc = ws.cell(row=r, column=1, value=header)
        lc.font = Font(name=FONT, size=10, bold=True)
        lc.fill = INPUT_FILL
        lc.border = BORDER
        vc = ws.cell(row=r, column=2, value=value)
        vc.font = Font(name=FONT, size=10, color="0000FF")
        vc.fill = INPUT_FILL
        vc.border = BORDER
        vc.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    r += 1
    note = ws.cell(row=r, column=1, value="Note")
    note.font = Font(name=FONT, size=10, bold=True)
    nv = ws.cell(row=r, column=2, value=(
        "Control content is a professional starting point, not legal advice. Tailor it to your entity, "
        "jurisdiction and processing profile before use on an engagement — in particular the DPDP "
        "Significant Data Fiduciary designation, which determines whether several controls apply at all."
    ))
    nv.font = Font(name=FONT, size=10, italic=True)
    nv.alignment = Alignment(wrap_text=True, vertical="top")


def build_rcm(ws, controls):
    # Row 1: group banner. Row 2: column headers. Data starts row 3.
    col_of = {}
    start = None
    current = None
    for idx, (field, header, width, group) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        col_of[field] = letter

        if group != current:
            if current is not None:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=idx - 1)
            current, start = group, idx
        gc = ws.cell(row=1, column=idx, value=group if idx == start else None)
        gc.fill = GROUP_FILLS[group]
        gc.font = Font(name=FONT, size=10, bold=True, color=INK)
        gc.alignment = Alignment(horizontal="center", vertical="center")
        gc.border = BORDER

        hc = ws.cell(row=2, column=idx, value=header)
        hc.fill = HEADER_FILL
        hc.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        hc.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        hc.border = BORDER
    ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=len(COLUMNS))
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 34

    for r, control in enumerate(controls, start=3):
        for idx, (field, _, _, group) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=r, column=idx, value=cell_value(control, field))
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDER
            if group == FIELDWORK_GROUP:
                c.fill = INPUT_FILL
            if field == "inherent_risk":
                risk = control.get("inherent_risk")
                if risk in RISK_FILLS:
                    c.fill = RISK_FILLS[risk]
                    c.font = RISK_FONTS[risk]
                c.alignment = Alignment(horizontal="center", vertical="top")
            if field == "key_control":
                c.alignment = Alignment(horizontal="center", vertical="top")
                if control.get("key_control"):
                    c.font = Font(name=FONT, size=10, bold=True, color=INK)

    last_row = 2 + len(controls)
    ws.freeze_panes = "B3"                       # keep Control ID and both header rows visible
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{last_row}"

    for field, options in DROPDOWNS.items():
        letter = col_of[field]
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Invalid value",
            error="Choose one of: " + ", ".join(options),
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}{last_row}")

    return col_of, last_row


def build_summary(ws, col_of, last_row, controls):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 38
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 15

    fw = col_of["framework"]
    dom = col_of["domain"]
    risk = col_of["inherent_risk"]
    key = col_of["key_control"]
    result = col_of["test_result"]

    rng = lambda letter: f"RCM!${letter}$3:${letter}${last_row}"

    ws["A1"] = "Summary"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=INK)
    ws["A2"] = "Counts update automatically as the RCM sheet is completed. Do not overwrite these cells."
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="6B7688")

    def header_row(r, first, rest):
        ws.cell(row=r, column=1, value=first).font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
        ws.cell(row=r, column=1).fill = HEADER_FILL
        for i, h in enumerate(rest, start=2):
            c = ws.cell(row=r, column=i, value=h)
            c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center")

    def label(r, text):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=10)
        c.border = BORDER
        return c

    def formula(r, col, f):
        c = ws.cell(row=r, column=col, value=f)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
        return c

    # By framework
    r = 4
    header_row(r, "By framework", ["Controls", "Key", "High risk", "Tested"])
    frameworks = sorted({c["framework"] for c in controls})
    for name in frameworks:
        r += 1
        label(r, name)
        formula(r, 2, f'=COUNTIF({rng(fw)},$A{r})')
        formula(r, 3, f'=COUNTIFS({rng(fw)},$A{r},{rng(key)},"Yes")')
        formula(r, 4, f'=COUNTIFS({rng(fw)},$A{r},{rng(risk)},"High")')
        formula(r, 5, f'=COUNTIFS({rng(fw)},$A{r},{rng(result)},"<>Not Tested")-COUNTIFS({rng(fw)},$A{r},{rng(result)},"")')
    r += 1
    tc = label(r, "Total")
    tc.font = Font(name=FONT, size=10, bold=True)
    for col in range(2, 6):
        c = formula(r, col, f"=SUM({get_column_letter(col)}5:{get_column_letter(col)}{r-1})")
        c.font = Font(name=FONT, size=10, bold=True)

    # By domain
    r += 3
    header_row(r, "By domain", ["Controls", "Key", "High risk", "Tested"])
    domains = sorted({c["domain"] for c in controls})
    for name in domains:
        r += 1
        label(r, name)
        formula(r, 2, f'=COUNTIF({rng(dom)},$A{r})')
        formula(r, 3, f'=COUNTIFS({rng(dom)},$A{r},{rng(key)},"Yes")')
        formula(r, 4, f'=COUNTIFS({rng(dom)},$A{r},{rng(risk)},"High")')
        formula(r, 5, f'=COUNTIFS({rng(dom)},$A{r},{rng(result)},"<>Not Tested")-COUNTIFS({rng(dom)},$A{r},{rng(result)},"")')

    # Test results
    r += 3
    header_row(r, "Testing status", ["Controls", "% of total"])
    for name in ["Not Tested", "Effective", "Effective with Exceptions", "Ineffective"]:
        r += 1
        label(r, name)
        formula(r, 2, f'=COUNTIF({rng(result)},$A{r})')
        pc = formula(r, 3, f'=IFERROR(B{r}/COUNTA({rng(fw)}),0)')
        pc.number_format = "0.0%"


def main():
    controls = load_controls()
    wb = Workbook()

    legend = wb.active
    legend.title = "Legend"
    build_legend(legend, len(controls), "data/seed/*.seed.js via scripts/build-library.js")

    rcm = wb.create_sheet("RCM")
    col_of, last_row = build_rcm(rcm, controls)

    summary = wb.create_sheet("Summary")
    build_summary(summary, col_of, last_row, controls)

    wb.save(OUT)
    print(f"Wrote {OUT} — {len(controls)} controls, rows 3 to {last_row}")


if __name__ == "__main__":
    main()
